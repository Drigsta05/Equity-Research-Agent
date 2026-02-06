"""DOCX report builder tool using python-docx."""

from __future__ import annotations

import json
import logging
from pathlib import Path

logger = logging.getLogger(__name__)

OUTPUT_DIR = Path("output")

DOCX_TOOL = {
    "name": "create_docx",
    "description": (
        "Create a professional equity research report as a .docx file. "
        "Provide the report content as structured sections. Each section has a "
        "heading, body text, and optional tables. The tool handles formatting."
    ),
    "input_schema": {
        "type": "object",
        "properties": {
            "filename": {
                "type": "string",
                "description": "Output filename, e.g., 'AAPL_report.docx'",
            },
            "title": {
                "type": "string",
                "description": "Report title, e.g., 'Apple Inc. (AAPL) — Equity Research Report'",
            },
            "subtitle": {
                "type": "string",
                "description": "Subtitle line, e.g., 'Initiating Coverage | Buy | Target: $250'",
            },
            "sections": {
                "type": "array",
                "description": "Ordered list of report sections.",
                "items": {
                    "type": "object",
                    "properties": {
                        "heading": {
                            "type": "string",
                            "description": "Section heading",
                        },
                        "level": {
                            "type": "integer",
                            "description": "Heading level: 1 for major sections, 2 for subsections",
                            "default": 1,
                        },
                        "body": {
                            "type": "string",
                            "description": (
                                "Section body text. Use \\n for paragraphs. "
                                "Use **bold** for emphasis. Use bullet points with '• ' prefix."
                            ),
                        },
                        "table": {
                            "type": "object",
                            "description": "Optional table for this section.",
                            "properties": {
                                "headers": {
                                    "type": "array",
                                    "items": {"type": "string"},
                                },
                                "rows": {
                                    "type": "array",
                                    "items": {
                                        "type": "array",
                                        "items": {"type": "string"},
                                    },
                                },
                            },
                        },
                    },
                    "required": ["heading", "body"],
                },
            },
        },
        "required": ["filename", "title", "sections"],
    },
}


def execute_create_docx(
    filename: str,
    title: str,
    sections: list[dict],
    subtitle: str = "",
) -> str:
    """Create a .docx report from structured sections."""
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT

        doc = Document()

        # Set default font
        style = doc.styles["Normal"]
        font = style.font
        font.name = "Calibri"
        font.size = Pt(11)
        font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        # Title
        title_para = doc.add_heading(title, level=0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Subtitle
        if subtitle:
            sub_para = doc.add_paragraph()
            sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = sub_para.add_run(subtitle)
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0x00, 0x55, 0x99)
            run.bold = True

        doc.add_paragraph("")  # spacing

        # Sections
        for section in sections:
            heading = section.get("heading", "")
            level = section.get("level", 1)
            body = section.get("body", "")
            table_data = section.get("table")

            # Add heading
            doc.add_heading(heading, level=min(level, 3))

            # Add body paragraphs
            for paragraph_text in body.split("\n"):
                paragraph_text = paragraph_text.strip()
                if not paragraph_text:
                    continue

                if paragraph_text.startswith("• ") or paragraph_text.startswith("- "):
                    # Bullet point
                    p = doc.add_paragraph(paragraph_text[2:], style="List Bullet")
                else:
                    p = doc.add_paragraph()
                    # Handle **bold** markers
                    parts = paragraph_text.split("**")
                    for i, part in enumerate(parts):
                        if part:
                            run = p.add_run(part)
                            if i % 2 == 1:  # odd index = between ** markers = bold
                                run.bold = True

            # Add table if present
            if table_data:
                headers = table_data.get("headers", [])
                rows = table_data.get("rows", [])

                if headers:
                    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
                    table.style = "Light Grid Accent 1"
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER

                    # Header row
                    for i, header in enumerate(headers):
                        cell = table.rows[0].cells[i]
                        cell.text = header
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.bold = True

                    # Data rows
                    for row_idx, row_data in enumerate(rows):
                        for col_idx, cell_text in enumerate(row_data):
                            if col_idx < len(headers):
                                table.rows[row_idx + 1].cells[col_idx].text = str(cell_text)

                    doc.add_paragraph("")  # spacing after table

        # Save
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        filepath = OUTPUT_DIR / filename
        doc.save(str(filepath))

        return json.dumps({
            "success": True,
            "filename": filename,
            "sections_count": len(sections),
            "path": str(filepath),
        })

    except ImportError:
        return json.dumps({
            "error": "python-docx not installed. Run: pip install python-docx"
        })
    except Exception as e:
        logger.exception("DOCX creation failed")
        return json.dumps({"error": f"DOCX creation failed: {str(e)}"})
